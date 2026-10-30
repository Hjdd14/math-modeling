"""design_board 决策板的 xlsx 渲染与回读。
把 design_board.json 渲染成 7 个工作表的 Excel 决策板，让用户不碰 JSON 也能确认
建模方案；用户改完后，再由 AI 用 ``read_submission_xlsx`` 读回成 ingest 可直接
使用的提交对象。
约定：黄底（FFF2CC）可改、灰底（F2F2F2）只读，01_待决策 已加工作表保护；确认码由
决策板哈希派生，必须用户本人在终端输入，AI 不得代填。只依赖 openpyxl。
"""

from __future__ import annotations

import json
import secrets
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAS_OPENPYXL = False

SHEET_OVERVIEW = "00_总览"
SHEET_DECIDE = "01_待决策"
SHEET_WEIGHTS = "02_权重"
SHEET_SCORES = "03_打分"
SHEET_SOURCES = "04_来源"
SHEET_COMMENTS = "05_批注"
SHEET_README = "06_说明"
#: B 闸门多一张数据对比表（可比层 + 专属层 + baseline 参照行），说明表顺延为 07。
SHEET_DATA = "06_数据对比"
SHEET_README_GATE_B = "07_说明"
TIER_LABELS = {"recommended": "🟢 推荐", "alternative": "⚪ 备选", "rejected": "⛔ 不推荐"}
WEIGHT_DIMENSIONS = [
    ("problem_fit", "问题契合"), ("interpretability", "可解释性"), ("computability", "可计算性"),
    ("data_requirement", "数据需求"), ("robustness", "稳健性"), ("validation_difficulty", "验证难度"),
    ("paper_presentation", "论文呈现"),
]
WEIGHT_NOTES = {
    "problem_fit": "方案对题目核心问题与约束的覆盖程度，越高越贴合",
    "interpretability": "结果能否用机理或业务语言讲清，评委能否看懂",
    "computability": "赛时算力与数据下能否跑通，含规模、收敛与数值稳定",
    "data_requirement": "对附件数据的依赖与清洗成本，越高表示越省数据",
    "robustness": "参数扰动、异常值与缺失数据下结论的稳定程度",
    "validation_difficulty": "能否给出可复核的检验口径，越高表示越容易验证",
    "paper_presentation": "图表、公式与行文表达的发挥空间",
}
READONLY_MESSAGE = "渲染/回读 xlsx 需要 openpyxl：pip install openpyxl"
PRIMARY, EDITABLE, GREY, GOOD, WARN = "1F4E79", "FFF2CC", "F2F2F2", "E8F5E9", "C62828"

if HAS_OPENPYXL:
    _THIN = Side(style="thin", color="D9D9D9")
    BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
    FILL_GREY, FILL_EDIT = PatternFill("solid", fgColor=GREY), PatternFill("solid", fgColor=EDITABLE)
    FILL_GOOD, FILL_BAD = PatternFill("solid", fgColor=GOOD), PatternFill("solid", fgColor="FFC7CE")
    HEADER_FONT, HEADER_FILL = Font(bold=True, color="FFFFFF", size=11), PatternFill("solid", fgColor=PRIMARY)
    BODY_FONT, BOLD_FONT = Font(size=10), Font(size=10, bold=True)
    WARN_FONT, BODY_ALIGN = Font(size=10, color=WARN), Alignment(wrap_text=True, vertical="top")

def _resolve_gate(board: dict, gate: str | None) -> str:
    """渲染用闸门：板自报优先，其次显式参数，兜底为 A（旧板=legacy A 版式）。"""
    declared = board.get("gate")
    if isinstance(declared, str) and declared.strip().lower() in ("a", "b"):
        return declared.strip().lower()
    if isinstance(gate, str) and gate.strip().lower() == "b":
        return "b"
    return "a"


def _verdict_text(option: dict, candidate: dict) -> str:
    """B 板的建议结论（中文），与 HTML 视图同一口径。"""
    raw = str(option.get("verdict") or candidate.get("verdict") or "").strip().lower()
    if raw in ("recommended", "recommend"):
        return "🟢 建议采用"
    if raw in ("not_recommended", "not-recommended", "rejected", "reject"):
        return "⛔ 不建议"
    return "⚪ 未表态"


def _advice_of(board: dict, question: dict) -> dict:
    """某题的建议块：题内 advice 优先，其次板级 advice。"""
    advice = question.get("advice")
    if isinstance(advice, dict) and advice:
        return advice
    board_advice = board.get("advice")
    return board_advice if isinstance(board_advice, dict) else {}


def _candidate_index(advice: dict) -> dict[str, dict]:
    """advice.candidates 按 option_id 建索引。"""
    index: dict[str, dict] = {}
    for item in advice.get("candidates") or []:
        if isinstance(item, dict):
            option_id = _text(item.get("option_id") or item.get("id")).strip()
            if option_id:
                index[option_id] = item
    return index


def _reason_text(option: dict, candidate: dict) -> str:
    """一句话理由：建议项取 pros[0]，不建议项取 why_not[0]。"""
    verdict = str(option.get("verdict") or candidate.get("verdict") or "").strip().lower()
    key = "pros" if verdict in ("recommended", "recommend") else "why_not"
    for item in candidate.get(key) or []:
        text = _text(item).strip()
        if text:
            return text
    return ""


def render_xlsx(board: dict, path, *, gate: str | None = None) -> str:
    """渲染 design_board 为 xlsx，返回 str(path)。

    A 板保持改造前的 7 张工作表与列序（回归锚）；B 板追加候选分数列与
    「06_数据对比」表，说明表顺延为 07_说明。
    """
    _require_openpyxl()
    mode = _resolve_gate(board, gate)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    questions = [q for q in board.get("questions") or [] if isinstance(q, dict)]
    now_text = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
    wb = Workbook()
    wb.remove(wb.active)
    _overview(wb.create_sheet(SHEET_OVERVIEW), board, questions, now_text)
    _decide(wb.create_sheet(SHEET_DECIDE), questions, board=board, gate=mode)
    _weights(wb.create_sheet(SHEET_WEIGHTS))
    _scores(wb.create_sheet(SHEET_SCORES), questions)
    _sources(wb.create_sheet(SHEET_SOURCES), questions)
    _comments(wb.create_sheet(SHEET_COMMENTS), questions)
    if mode == "b":
        _data_comparison(wb.create_sheet(SHEET_DATA), board, questions)
        _readme(wb.create_sheet(SHEET_README_GATE_B), gate=mode, filename=target.name)
    else:
        _readme(wb.create_sheet(SHEET_README))
    wb.active = 0
    wb.save(target)
    return str(target)

def read_submission_xlsx(path, board: dict) -> dict:
    """读回用户改好的 xlsx，返回 ingest 可用的提交对象。"""
    _require_openpyxl()
    warnings: list[str] = []
    wb = load_workbook(Path(path), data_only=True)
    try:
        index = _question_index(board)
        current = {qid: _text(q.get("user_comment")) for qid, (q, _) in index.items()}
        decisions, comments = _read_decisions(wb, index, current, warnings)
        weights, scores = _read_weights(wb, warnings), _read_scores(wb, index, warnings)
        edited = any(item["option_id"] != _text(index[item["question_id"]][0].get("user_choice")) for item in decisions)
    finally:
        wb.close()
    return {
        "board_id": board.get("board_id"), "at": datetime.now().astimezone().isoformat(),
        "receipt_origin": "xlsx-upload", "nonce": secrets.token_hex(8), "user_edited": edited,
        "decisions": decisions, "comments": comments, "weights": weights, "scores": scores,
        "warnings": warnings,
    }

def _overview(ws, board: dict, questions: list[dict], now_text: str) -> None:
    """00_总览：只读灰底。"""
    ws.sheet_view.showGridLines = False
    for row, height in ((1, 36), (2, 26), (3, 32)):
        _merge_row(ws, row, 1, 4)
        ws.row_dimensions[row].height = height
    title = ws.cell(row=1, column=1, value="建模方案决策板")
    title.font, title.alignment = Font(bold=True, size=18, color=PRIMARY), Alignment(horizontal="center", vertical="center")
    code = ws.cell(row=2, column=1, value=f"确认码：{_text(board.get('confirm_code')) or '—'}")
    code.font, code.alignment = Font(bold=True, size=14, color="2E7D32"), Alignment(horizontal="left", vertical="center")
    hint = ws.cell(row=3, column=1, value="改完 01_待决策 的当前选择后保存，交给 AI 用 ingest 回读；确认码需你本人在终端亲自输入")
    hint.alignment = Alignment(wrap_text=True, vertical="center")
    confirmed = _text(board.get("confirmed_at"))
    fields = [("board_id", _text(board.get("board_id")) or "—"), ("language", _text(board.get("language")) or "—"),
              ("revision", _text(board.get("revision")) or "—"), ("状态", f"已确认（{confirmed}）" if confirmed else "待确认")]
    for offset, (label, value) in enumerate(fields):
        _merge_row(ws, 4 + offset, 2, 4)
        _write_row(ws, 4 + offset, [label, value])
        ws.cell(row=4 + offset, column=1).font = BOLD_FONT
    row = 5 + len(fields)
    for question in questions:
        _write_row(ws, row, [_text(question.get("id")) or "—", _text(question.get("goal"))[:60] or "—",
                             _text(question.get("user_choice")) or "（未选）", len(question.get("options") or [])])
        row += 1
    row += 1
    _merge_row(ws, row, 2, 4)
    _write_row(ws, row, ["生成时间", now_text])
    ws.cell(row=row, column=1).font = BOLD_FONT
    _autofit(ws)

def _decide(ws, questions: list[dict], *, board: dict | None = None, gate: str = "a") -> None:
    """01_待决策：黄底可编辑，只有「当前选择」「批注」两列解锁。

    A 板列序与改造前逐字一致；B 板把「推荐度/理由/复杂度/风险/试算」换成
    「建议/总分/范式/来源视角/推荐或不建议理由」，其余保持同构，回读逻辑不变。
    """
    if gate == "b":
        headers = [
            "题号", "目标", "选项ID", "选项名", "建议", "总分", "范式", "来源视角",
            "推荐或不建议理由", "当前选择", "批注",
        ]
    else:
        headers = [
            "题号", "目标", "选项ID", "选项名", "推荐度", "一句话理由", "复杂度",
            "主要风险", "试算结果", "当前选择", "批注",
        ]
    _header_row(ws, 1, headers)
    choice_col = headers.index("当前选择") + 1
    comment_col = headers.index("批注") + 1
    verdict_col = headers.index("建议" if gate == "b" else "推荐度") + 1
    row, option_ids = 2, []
    for question in questions:
        qid, first = _text(question.get("id")), row
        choice, comment = _text(question.get("user_choice")), _text(question.get("user_comment"))
        advice = _advice_of(board or {}, question)
        candidates = _candidate_index(advice)
        for option in [o for o in question.get("options") or [] if isinstance(o, dict)]:
            oid, tier = _text(option.get("id")), _text(option.get("tier"))
            if oid and oid not in option_ids: option_ids.append(oid)
            candidate = candidates.get(oid, {})
            if gate == "b":
                score = option.get("score") if option.get("score") is not None else candidate.get("score")
                values = [
                    qid, _text(question.get("goal")), oid, _text(option.get("name")),
                    _verdict_text(option, candidate),
                    float(score) if isinstance(score, (int, float)) and not isinstance(score, bool) else None,
                    _text(candidate.get("paradigm_id") or option.get("paradigm_id")),
                    _text(candidate.get("origin") or option.get("origin")),
                    _reason_text(option, candidate),
                    choice if row == first else "", comment if row == first else "",
                ]
            else:
                values = [
                    qid, _text(question.get("goal")), oid, _text(option.get("name")),
                    TIER_LABELS.get(tier, tier or "—"),
                    _text(option.get("rationale")), _text(option.get("complexity")), _text(option.get("risk")),
                    _trial_text(option), choice if row == first else "", comment if row == first else "",
                ]
            _write_row(ws, row, values, editable=(choice_col, comment_col))
            verdict = str(option.get("verdict") or "").strip().lower() if gate == "b" else tier
            if verdict in ("recommended", "recommend"):
                ws.cell(row=row, column=verdict_col).fill = FILL_GOOD
            elif verdict in ("rejected", "not_recommended", "not-recommended"):
                ws.cell(row=row, column=verdict_col).font = WARN_FONT
            row += 1
    last, formula = row - 1, '"' + ",".join(option_ids) + '"'
    if option_ids and last >= 2 and len(formula) <= 255:
        dv = DataValidation(type="list", formula1=formula, allow_blank=False)
        dv.promptTitle, dv.prompt = "当前选择", "从下拉列表选择本题最终采用的选项ID"
        dv.errorTitle, dv.error = "选项ID 不合法", "请从下拉列表里选择本题的选项ID"
        dv.add(f"{get_column_letter(choice_col)}2:{get_column_letter(choice_col)}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = False
    ws.protection.sheet = True
    _autofit(ws)


def _data_comparison(ws, board: dict, questions: list[dict]) -> None:
    """06_数据对比（B 板）：每个 slice 一组行，可比层可横比，专属层单列，baseline 收尾。"""
    rows: list[list] = []
    columns: list[str] = []
    for question in questions:
        advice = _advice_of(board, question)
        comparison = advice.get("data_comparison") if isinstance(advice.get("data_comparison"), dict) else {}
        for name in comparison.get("comparable_columns") or []:
            text = _text(name).strip()
            if text and text not in columns:
                columns.append(text)
        rows.append((question, comparison))
    headers = ["题号", "切片", "候选", "回溯", *columns, "专属指标（own）", "备注"]
    _header_row(ws, 1, headers)
    row = 2
    for question, comparison in rows:
        qid = _text(question.get("id"))
        data_rows = [item for item in comparison.get("rows") or [] if isinstance(item, dict)]
        baseline = comparison.get("baseline_row") if isinstance(comparison.get("baseline_row"), dict) else None
        for item in data_rows:
            comparable = item.get("comparable") if isinstance(item.get("comparable"), dict) else {}
            own = item.get("own") if isinstance(item.get("own"), dict) else {}
            _write_row(ws, row, [
                qid, _text(item.get("slice_id")), _text(item.get("option_id")), _text(item.get("trials_ref")),
                *[_text(comparable.get(column)) for column in columns],
                "、".join(f"{key}={_text(value)}" for key, value in own.items()),
                "",
            ])
            row += 1
        if baseline:
            comparable = baseline.get("comparable") if isinstance(baseline.get("comparable"), dict) else {}
            _write_row(ws, row, [
                qid, _text(baseline.get("slice_id")), _text(baseline.get("option_id")) or "baseline", "",
                *[_text(comparable.get(column)) for column in columns],
                "", "参照行（不占候选名额）",
            ])
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).font = BOLD_FONT
            row += 1
    if row == 2:
        _write_row(ws, row, ["（本板没有 data_comparison：先跑 perspective_trial_runner 再生成建议）"])
    _autofit(ws)

def _weights(ws) -> None:
    """02_权重：默认等权 1/7，合计不等于 1 标红。

    前 6 维写 `=1/7` 公式、最后一维写 `=1-SUM(前 6 维)`，这样合计**恒为 1**，
    「合计 ≠ 1」的红色条件格式只会在用户真正改坏权重时亮起。
    若 7 维都写 `round(1/7, 4)=0.1429`，合计是 1.0003，表一打开就是红的——
    那会让告警失去信号价值（和 post_confirm_edit 踩过的坑同类）。
    """
    _header_row(ws, 1, ["维度", "权重", "说明"])
    last_row = 2 + len(WEIGHT_DIMENSIONS) - 1
    for offset, (key, label) in enumerate(WEIGHT_DIMENSIONS):
        row = 2 + offset
        if row == last_row:
            value: Any = f"=1-SUM(B2:B{last_row - 1})"
        else:
            value = "=1/7"
        _write_row(ws, row, [label, value, f"{WEIGHT_NOTES[key]}（{key}）"], editable=(2,))
        ws.cell(row=row, column=2).number_format = "0.0000"
    total = 2 + len(WEIGHT_DIMENSIONS)
    _write_row(ws, total, ["合计", f"=SUM(B2:B{total - 1})", "权重合计必须等于 1"])
    ws.cell(row=total, column=2).number_format = "0.0000"
    ws.conditional_formatting.add(f"B{total}", CellIsRule(operator="notEqual", formula=["1"], fill=FILL_BAD))
    _autofit(ws)

def _scores(ws, questions: list[dict]) -> None:
    """03_打分：D:J 填 0-100，K 列按 02_权重 加权。"""
    headers = ["题号", "选项ID", "选项名", "契合", "可解释", "可计算", "数据", "稳健", "验证难度", "论文呈现", "加权总分"]
    _header_row(ws, 1, headers)
    row = 2
    for question in questions:
        qid = _text(question.get("id"))
        for option in [o for o in question.get("options") or [] if isinstance(o, dict)]:
            terms = [f"{get_column_letter(4 + i)}{row}*'02_权重'!$B${2 + i}" for i in range(7)]
            values = [qid, _text(option.get("id")), _text(option.get("name"))] + [None] * 7 + ["=" + "+".join(terms)]
            _write_row(ws, row, values, editable=tuple(range(4, 11)))
            ws.cell(row=row, column=11).number_format = "0.00"
            row += 1
    last = row - 1
    if last >= 2:
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
        dv.errorTitle, dv.error = "分数必须是整数", "只能填 0 到 100 之间的整数"
        dv.add(f"D2:J{last}")
        ws.add_data_validation(dv)
        ws.conditional_formatting.add(f"K2:K{last}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
            mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    _autofit(ws)

def _sources(ws, questions: list[dict]) -> None:
    """04_来源：只读，未验证用红字提示。"""
    _header_row(ws, 1, ["选项ID", "来源类型", "来源引用", "可信度", "备注"])
    row = 2
    for question in questions:
        for option in [o for o in question.get("options") or [] if isinstance(o, dict)]:
            ref = option.get("source_ref") if isinstance(option.get("source_ref"), dict) else {}
            _write_row(ws, row, [
                _text(option.get("id")) or "—", _text(ref.get("type")) or "—", _text(ref.get("ref")) or "—",
                "已验证" if ref.get("verified") is True else "未验证", _text(ref.get("note")),
            ])
            if ref.get("verified") is not True: ws.cell(row=row, column=4).font = WARN_FONT
            row += 1
    _autofit(ws)

def _comments(ws, questions: list[dict]) -> None:
    """05_批注：每 option 一行，末尾留 2 行空行。"""
    _header_row(ws, 1, ["题号", "选项ID", "批注"])
    row = 2
    for question in questions:
        qid, comment, first = _text(question.get("id")), _text(question.get("user_comment")), row
        for option in [o for o in question.get("options") or [] if isinstance(o, dict)]:
            _write_row(ws, row, [qid, _text(option.get("id")), comment if row == first else ""], editable=(3,))
            row += 1
    for _ in range(2):
        _write_row(ws, row, [None, None, None], editable=(3,))
        row += 1
    _autofit(ws)

def _readme(ws, *, gate: str = "a", filename: str = "design_choices.xlsx") -> None:
    """说明表：A 板 06_说明（9 行），B 板 07_说明（多一段闸门与文件名说明）。"""
    command = (
        f"python tools/design_board.py ingest --workspace . --gate b --from {filename}"
        if gate == "b"
        else f"python tools/design_board.py ingest --workspace . --from {filename}"
    )
    lines = [
        "xlsx 决策板使用说明",
        "1) 怎么用：00_总览 看整体结论；01_待决策 逐题确认当前选择；02_权重、03_打分 可调权重与打分并看加权总分。",
        "2) 哪些格子能改：只有黄底可改（01 的 J/K 列、02 的权重 B2:B8、03 的 D:J 分数、05 的 C 列批注）；灰底只读，01 已加保护。",
        f"3) 改完怎么交回：另存为 {filename}，运行 {command}",
        "4) 确认码含义：00_总览 A2 的 4 位确认码由决策板哈希派生，必须你本人在终端亲自输入，AI 不得代填或代为确认。",
        "5) 什么情况需重新确认：AI 改了候选方案内容（目标、选项、理由、风险、试算等）会让 board_sha256 变、确认码随之变，必须重新确认。",
        "6) 什么情况不需重新确认：你只改自己的选择或批注，hash 未变，改完保存交回即可，无需重新确认。",
        "7) 校验：02_权重 的合计应为 1（B9 不等于 1 会标红，请微调到恰好为 1）；03_打分 的 D:J 只能填 0-100 的整数。",
        "8) 回读：AI 用 read_submission_xlsx 解析本文件，改动会以 receipt_origin=xlsx-upload 记录并交给 ingest。",
    ]
    if gate == "b":
        lines.extend([
            "9) 闸门说明：本表属于 B 闸门（方案择定），01_待决策 的「建议/总分/范式/来源视角」是建议 Agent 的结论，不是评委审查。",
            "10) 数据对比：06_数据对比 的可比层（comparable）可以横向比较；专属指标（own）只在自己这个候选的不同切片之间可比；baseline 行是参照，不占候选名额。",
            "11) 回执位置：B 板的回执写入 design_rounds_gate_b/，与 A 板的 design_rounds/ 互不干扰。",
        ])
    for row, line in enumerate(lines, start=1):
        cell = ws.cell(row=row, column=1, value=line)
        cell.font, cell.fill, cell.border, cell.alignment = (BOLD_FONT if row == 1 else BODY_FONT), FILL_GREY, BORDER, BODY_ALIGN
        ws.row_dimensions[row].height = 26 if row == 1 else 42
    ws.column_dimensions["A"].width = 60

def _read_decisions(wb, index: dict, current: dict, warnings: list[str]) -> tuple[list[dict], list[dict]]:
    """01 逐行回读选择与批注，再合并 05 的批注；非法值只记 warning 不抛错。"""
    ws = _open_sheet(wb, SHEET_DECIDE)
    header_row, columns = _locate(ws, ("题号", "当前选择"))
    try:
        comment_col = _locate(ws, ("批注",))[1]["批注"]
    except ValueError:
        comment_col = 0
        warnings.append("01_待决策 缺少「批注」列，本次不回读批注")
    picked: dict[str, list[str]] = {}
    notes: dict[str, str] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        qid = _text(ws.cell(row=row, column=columns["题号"]).value).strip()
        if not qid: continue
        if qid not in index:
            warnings.append(f"01_待决策 第 {row} 行题号 {qid} 不在决策板中，已跳过")
            continue
        choice = _text(ws.cell(row=row, column=columns["当前选择"]).value).strip()
        if choice in index[qid][1]:
            picked.setdefault(qid, []).append(choice)
        elif choice:
            warnings.append(f"01_待决策 第 {row} 行 {qid} 的当前选择 {choice} 不是合法选项，已跳过")
        note = _text(ws.cell(row=row, column=comment_col).value).strip() if comment_col else ""
        if note: notes.setdefault(qid, note)
    try:  # 05_批注 可单独编辑：只有它给出改过的内容时才覆盖 01 的批注
        ws5 = _open_sheet(wb, SHEET_COMMENTS)
        h5, c5 = _locate(ws5, ("题号", "批注"))
        for row in range(h5 + 1, ws5.max_row + 1):
            qid, text = (_text(ws5.cell(row=row, column=c5[key]).value).strip() for key in ("题号", "批注"))
            if qid in index and text and (qid not in notes or notes[qid] == current.get(qid)): notes[qid] = text
    except ValueError:
        pass
    decisions = []
    for qid in index:
        values = picked.get(qid) or []
        if not values: continue
        if len(set(values)) > 1:
            warnings.append(f"{qid} 出现多个不同的当前选择，取最后一个：{values[-1]}")
        decisions.append({"question_id": qid, "option_id": values[-1]})
    return decisions, [{"question_id": qid, "text": notes[qid]} for qid in index if notes.get(qid)]

def _read_weights(wb, warnings: list[str]):
    try:
        ws = _open_sheet(wb, SHEET_WEIGHTS)
        header_row, columns = _locate(ws, ("维度", "权重"))
    except ValueError as exc:
        warnings.append(f"未能回读权重：{exc}")
        return None
    labels = {label: key for key, label in WEIGHT_DIMENSIONS}
    keys = {key for key, _ in WEIGHT_DIMENSIONS}
    weights: dict[str, float] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        name = _text(ws.cell(row=row, column=columns["维度"]).value).strip()
        if not name or name == "合计": continue
        key = name if name in keys else labels.get(name)
        value = ws.cell(row=row, column=columns["权重"]).value
        if not key:
            warnings.append(f"02_权重 第 {row} 行维度 {name} 无法识别，已忽略")
        elif isinstance(value, bool) or not isinstance(value, (int, float)):
            warnings.append(f"02_权重 {name} 的权重不是数字，已忽略")
        else:
            weights[key] = float(value)
    if not weights: return None
    missing = [label for key, label in WEIGHT_DIMENSIONS if key not in weights]
    if missing: warnings.append(f"02_权重 缺少维度：{'、'.join(missing)}")
    return weights

def _read_scores(wb, index: dict, warnings: list[str]):
    try:
        ws = _open_sheet(wb, SHEET_SCORES)
        header_row, columns = _locate(ws, ("题号", "选项ID", "加权总分"))
    except ValueError as exc:
        warnings.append(f"未能回读分数：{exc}")
        return None
    rows: list[dict] = []
    uncalculated = 0
    for row in range(header_row + 1, ws.max_row + 1):
        qid = _text(ws.cell(row=row, column=columns["题号"]).value).strip()
        oid = _text(ws.cell(row=row, column=columns["选项ID"]).value).strip()
        if not qid or not oid: continue
        total = ws.cell(row=row, column=columns["加权总分"]).value
        if qid not in index or oid not in index[qid][1]:
            warnings.append(f"03_打分 第 {row} 行 {qid}/{oid} 不在决策板候选中，已跳过")
        elif isinstance(total, bool) or not isinstance(total, (int, float)):
            uncalculated += 1
        else:
            rows.append({"question_id": qid, "option_id": oid, "total": float(total)})
    if uncalculated:
        warnings.append(f"03_打分 有 {uncalculated} 行加权总分为空（公式未计算），已跳过；用 Excel 保存后可回读")
    return rows or None

def _require_openpyxl() -> None:
    if not HAS_OPENPYXL:
        raise RuntimeError(READONLY_MESSAGE)

def _write_row(ws, row: int, values: list, editable: tuple = ()) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = value
        cell.font, cell.alignment, cell.border = BODY_FONT, BODY_ALIGN, BORDER
        cell.fill = FILL_EDIT if col in editable else FILL_GREY
        cell.protection = Protection(locked=col not in editable)

def _header_row(ws, row: int, labels: list) -> None:
    ws.row_dimensions[row].height = 24
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def _merge_row(ws, row: int, first: int, last: int) -> None:
    ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    for col in range(first, last + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.border = FILL_GREY, BORDER

def _autofit(ws, minimum: int = 8, maximum: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is None or (isinstance(value, str) and value.startswith("=")): continue
            width = max(sum(2 if ord(ch) > 0x2E80 else 1 for ch in line) for line in str(value).split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), width)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = float(min(max(minimum, width + 2), maximum))

def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)

def _trial_text(option: dict) -> str:
    trial = option.get("mini_trial")
    if not isinstance(trial, dict):
        return "未试算"
    parts: list[str] = []
    output = trial.get("output")
    if isinstance(output, dict):
        parts.append(", ".join(f"{key}={_text(value)}" for key, value in output.items()))
    elif output not in (None, ""):
        parts.append(_text(output))
    seconds = trial.get("seconds")
    if isinstance(seconds, (int, float)) and not isinstance(seconds, bool):
        parts.append(f"{seconds * 1000:.2f}ms")
    return ", ".join(part for part in parts if part) or "未试算"

def _question_index(board: dict) -> dict[str, tuple[dict, list[str]]]:
    index: dict[str, tuple[dict, list[str]]] = {}
    for question in board.get("questions") or []:
        if not isinstance(question, dict):
            continue
        qid = _text(question.get("id")).strip()
        if qid and qid not in index:
            index[qid] = (question, [str(o.get("id")) for o in question.get("options") or [] if isinstance(o, dict) and o.get("id")])
    return index

def _open_sheet(wb, name: str):
    if name not in wb.sheetnames:
        raise ValueError(f"xlsx 缺少工作表：{name}")
    return wb[name]

def _locate(ws, names: tuple) -> tuple[int, dict[str, int]]:
    for row in range(1, min(ws.max_row, 8) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, min(ws.max_column, 60) + 1):
            text = _text(ws.cell(row=row, column=col).value).strip()
            if text and text not in mapping: mapping[text] = col
        if all(name in mapping for name in names):
            return row, {name: mapping[name] for name in names}
    raise ValueError(f"{ws.title} 找不到表头列：{'、'.join(names)}")
