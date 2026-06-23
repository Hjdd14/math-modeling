#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""题目文件解析与 problem_brief.md 生成工具"""

import argparse
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

try:
    from tools.data_analyzer import check_quality, generate_stats, read_data_structure
except ImportError:
    from data_analyzer import check_quality, generate_stats, read_data_structure


def output(result: dict):
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


def error(message: str, code: int = 1):
    print(json.dumps({"error": message}, ensure_ascii=False), file=sys.stderr)
    sys.exit(code)


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("缺少 pypdf，请先安装 requirements.txt") from exc

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(part for part in parts if part).strip()


def _read_docx_with_zip(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    text_parts = []
    for node in root.iter():
        if node.tag.endswith("}t") and node.text:
            text_parts.append(node.text)
    return "\n".join(text_parts).strip()


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        return _read_docx_with_zip(path)

    try:
        document = Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    parts.append(" | ".join(row_text))
        return "\n".join(parts).strip()
    except Exception:
        return _read_docx_with_zip(path)


def extract_text(file_path: str) -> dict:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {file_path}")

    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        text = _read_text_file(path)
    elif ext == ".pdf":
        text = _read_pdf(path)
    elif ext == ".docx":
        text = _read_docx(path)
    else:
        raise ValueError(f"不支持的题目文件格式: {ext}")

    return {
        "file_name": path.name,
        "file_type": ext.lstrip("."),
        "text": text,
        "char_count": len(text),
        "line_count": len(text.splitlines()),
    }


def summarize_data_file(file_path: str) -> dict:
    return {
        "structure": read_data_structure(file_path),
        "stats": generate_stats(file_path),
        "quality": check_quality(file_path),
    }


def _first_lines(text: str, limit: int = 1200) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "\n\n[题目文本已截断，完整文本见原文件]"


def render_brief(problem_info: dict, data_summaries: list) -> str:
    lines = [
        "# 数学建模题目简报",
        "",
        "## 1. 题目原文",
        "",
        _first_lines(problem_info["text"]),
        "",
        "## 2. 求解目标",
        "",
        "- 待 Agent 根据题目原文提取并确认",
        "",
        "## 3. 已知条件与参数",
        "",
        "- 待 Agent 根据题目原文和附件数据提取",
        "",
        "## 4. 约束条件",
        "",
        "- 待 Agent 根据题目原文提取",
        "",
        "## 5. 附件数据",
        "",
        "### 5.1 附件清单",
        "",
        "| 文件名 | 工作表名 | 行数 | 列数 | 质量评分 |",
        "|--------|----------|------|------|----------|",
    ]

    for summary in data_summaries:
        structure = summary["structure"]
        quality = summary["quality"]
        for sheet in structure["sheets"]:
            lines.append(
                f"| {structure['file_name']} | {sheet['sheet_name']} | "
                f"{sheet['rows']} | {sheet['columns']} | {quality['quality_score']} |"
            )

    lines.extend([
        "",
        "### 5.2 数据结构与示例",
        "",
    ])

    for summary in data_summaries:
        structure = summary["structure"]
        for sheet in structure["sheets"]:
            cols = ", ".join(str(c) for c in sheet["column_names"])
            lines.extend([
                f"#### {structure['file_name']} / {sheet['sheet_name']}",
                "",
                f"- 行数：{sheet['rows']}",
                f"- 列数：{sheet['columns']}",
                f"- 字段：{cols}",
                "",
            ])

    lines.extend([
        "## 6. 输出要求",
        "",
        "- 待 Agent 根据题目原文提取",
        "",
        "## 7. 评分标准（如有）",
        "",
        "- 待 Agent 根据题目原文提取",
        "",
    ])
    return "\n".join(lines)


def build_brief(problem_file: str, data_files: list[str], out_path: str) -> dict:
    problem_info = extract_text(problem_file)
    data_summaries = [summarize_data_file(path) for path in data_files]
    content = render_brief(problem_info, data_summaries)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(content, encoding="utf-8")

    return {
        "brief_path": str(out),
        "problem": problem_info,
        "data_summaries": data_summaries,
        "written": True,
    }


def main():
    parser = argparse.ArgumentParser(description="题目文件解析与 problem_brief.md 生成工具")
    subparsers = parser.add_subparsers(dest="action", help="可用操作")

    p_extract = subparsers.add_parser("extract", help="提取题目文本")
    p_extract.add_argument("--file", required=True, help="PDF/DOCX/MD/TXT 文件路径")

    p_brief = subparsers.add_parser("brief", help="生成 problem_brief.md")
    p_brief.add_argument("--problem", required=True, help="题目文件路径")
    p_brief.add_argument("--data", nargs="*", default=[], help="附件数据文件路径")
    p_brief.add_argument("--out", required=True, help="输出 problem_brief.md 路径")

    args = parser.parse_args()
    if not args.action:
        parser.print_help()
        sys.exit(1)

    try:
        if args.action == "extract":
            output(extract_text(args.file))
        elif args.action == "brief":
            output(build_brief(args.problem, args.data, args.out))
    except Exception as exc:
        error(str(exc))


if __name__ == "__main__":
    main()
