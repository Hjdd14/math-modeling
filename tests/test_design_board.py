#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""design_board 决策板工具与 design_board_checker 检查器测试"""

import json
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pytest

from tools import design_board
from tools import design_board_checker


SKILL_DIR = Path(__file__).resolve().parents[1]
BOARD_TOOL = SKILL_DIR / "tools" / "design_board.py"
CONFIRMED_AT = "2026-06-28T12:00:00+00:00"
BOARD_ID = "b-2026-06-28-fixture"


# ---------------------------------------------------------------------------
# 夹具辅助
# ---------------------------------------------------------------------------


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def make_option(option_id, name, objective):
    return {
        "id": option_id,
        "name": name,
        "tier": "recommended" if option_id == "O1" else "alternative",
        "rationale": "约束线性、规模可控，可证明最优",
        "complexity": "O(n)",
        "dependencies": [],
        "risk": "规模翻倍时需要换求解器",
        "source_ref": {
            "type": "method_card",
            "ref": "OPT-01",
            "tier": "tier_a",
            "verified": True,
        },
        "mini_trial": {
            "sample": "前 20 行真实数据",
            "output": {"objective": objective, "feasible": True},
            "seconds": 0.01,
            "reproduce_cmd": "python solution.py",
        },
    }


def make_board(option_names=("整数规划", "贪心启发式"), user_choice="O1", no_change_rounds=2):
    """最小合法板：每个小问 2 个候选 + 导图 + 用户选择。"""
    return {
        "schema_version": 1,
        "board_id": BOARD_ID,
        "language": "python",
        "created_at": "2026-06-28T11:00:00+00:00",
        "revision": 2,
        "problem_summary": {
            "title": "生产计划优化",
            "questions": ["问题一"],
            "key_ambiguities": [],
        },
        "mindmap": {"format": "mermaid", "source": "graph LR\n  P[题目] --> Q1[问题一]"},
        "questions": [
            {
                "id": "Q1",
                "goal": "在产能与材料约束内最大化总利润",
                "options": [
                    make_option("O1", option_names[0], 12000),
                    make_option("O2", option_names[1], 11000),
                ],
                "user_choice": user_choice,
                "user_locked": False,
                "user_comment": "",
                "confirmed": False,
            }
        ],
        "pending_decisions": [],
        "no_change_rounds": no_change_rounds,
    }


def write_rounds(workspace, *, user_edited=(False, False)):
    """写两轮回执：changes 为空，checker 不会判成"确认后改动"。"""
    rounds_dir = workspace / design_board.ROUNDS_DIR
    rounds_dir.mkdir(parents=True, exist_ok=True)
    for index, edited in enumerate(user_edited, start=1):
        write_json(
            rounds_dir / f"round_{index:03d}.json",
            {
                "round": index,
                "at": f"2026-06-28T1{index}:00:00+00:00",
                "board_id": BOARD_ID,
                "board_sha256": "",
                "user_edited": edited,
                "changes": [],
                "receipt_origin": "browser-submit",
                "receipt_nonce": f"nonce-{index}",
            },
        )


def write_audit(workspace, candidate_names, final_name):
    write_json(
        workspace / "model_selection_audit.json",
        {
            "candidates": [
                {"name": name, "score": round(0.9 - index * 0.1, 2), "evidence": ["oracle"]}
                for index, name in enumerate(candidate_names)
            ],
            "final_selection": {"name": final_name, "score": 0.9, "evidence": ["optimality_gap=0"]},
        },
    )


def seed_board(workspace, board=None):
    board = board if board is not None else make_board()
    write_json(workspace / "design_board.json", board)
    return board


def build_confirmed_board(
    *,
    option_names=("整数规划", "贪心启发式"),
    user_choice="O1",
    confirmed_at=CONFIRMED_AT,
    no_change_rounds=2,
):
    """构造一份"已确认"的板：哈希与确认码都按权威实现真实计算。"""
    board = make_board(option_names=option_names, user_choice=user_choice, no_change_rounds=no_change_rounds)
    board["questions"][0]["confirmed"] = True
    board["confirmed"] = True
    board["confirmed_at"] = confirmed_at
    board["confirmed_by"] = "user"
    board["receipt_origin"] = "browser-submit"
    board["confirm_attested"] = True
    sha = design_board.compute_board_sha256(board)
    board["board_sha256"] = sha
    board["confirm_code"] = design_board.derive_confirm_code(sha)
    return board


def seed_confirmed_workspace(workspace, **kwargs):
    option_names = kwargs.pop("option_names", ("整数规划", "贪心启发式"))
    board = build_confirmed_board(option_names=option_names, **kwargs)
    write_json(workspace / "design_board.json", board)
    write_rounds(workspace)
    chosen_index = int(str(board["questions"][0]["user_choice"])[1:]) - 1
    write_audit(workspace, list(option_names), option_names[chosen_index])
    return board


def seed_ready_workspace(workspace):
    """满足全部收敛判据、可执行 confirm --apply 的工作区。"""
    board = make_board()
    board["confirmed"] = False
    write_json(workspace / "design_board.json", board)
    write_rounds(workspace)
    design_board.render_board(str(workspace))
    return read_json(workspace / "design_board.json")


def touch_at(path, iso_timestamp, content="print('ok')\n"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    stamp = datetime.fromisoformat(iso_timestamp).timestamp()
    os.utime(path, (stamp, stamp))


def issue_codes(result):
    return {issue["code"] for issue in result["issues"]}


def run_cli(*args):
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    return subprocess.run(
        [sys.executable, str(BOARD_TOOL), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(SKILL_DIR),
        env=env,
    )


# ---------------------------------------------------------------------------
# 1. render：回写真实 sha、派生确认码、渲染 Markdown
# ---------------------------------------------------------------------------


def test_render_writes_real_sha_confirm_code_and_markdown(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_board(workspace)
    write_rounds(workspace)

    result = design_board.render_board(str(workspace))

    board = read_json(workspace / "design_board.json")
    assert len(board["board_sha256"]) == 64
    assert all(char in "0123456789abcdef" for char in board["board_sha256"])
    assert board["confirm_code"] == board["board_sha256"][:4].upper()
    assert len(board["confirm_code"]) == 4
    assert board["board_sha256"] == design_board.compute_board_sha256(board)
    assert result["confirm_code"] == board["confirm_code"]

    markdown = (workspace / "design_board.md").read_text(encoding="utf-8")
    assert board["confirm_code"] in markdown
    assert "确认码" in markdown
    assert "生产计划优化" in markdown


def test_render_is_idempotent_for_content_hash(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_board(workspace)

    first = design_board.render_board(str(workspace))
    markdown_first = (workspace / "design_board.md").read_text(encoding="utf-8")
    second = design_board.render_board(str(workspace))
    markdown_second = (workspace / "design_board.md").read_text(encoding="utf-8")

    assert first["board_sha256"] == second["board_sha256"]
    assert first["confirm_code"] == second["confirm_code"]
    assert markdown_first == markdown_second


# ---------------------------------------------------------------------------
# 2/3. confirm：错码必须失败且不写字段；正确码成功并留痕
# ---------------------------------------------------------------------------


def test_confirm_rejects_wrong_code_without_writing_confirmation_fields(tmp_path, monkeypatch):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    original = seed_ready_workspace(workspace)
    monkeypatch.setenv(design_board.CONFIRM_CODE_ENV, "WRONG")

    with pytest.raises(ValueError) as excinfo:
        design_board.confirm_board(str(workspace), apply=True)
    assert "确认码不正确" in str(excinfo.value)

    board = read_json(workspace / "design_board.json")
    assert board["confirmed"] is False
    assert board.get("confirmed_at") in (None, original.get("confirmed_at"))
    assert "confirm_attested" not in board or board["confirm_attested"] is not True
    assert board["questions"][0]["confirmed"] is False


def test_confirm_with_correct_code_attests_user_choice(tmp_path, monkeypatch):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seeded = seed_ready_workspace(workspace)
    monkeypatch.setenv(design_board.CONFIRM_CODE_ENV, seeded["confirm_code"])

    result = design_board.confirm_board(str(workspace), apply=True)

    assert result["confirmed"] is True
    assert result["confirmed_by"] == "ci"
    assert result["confirm_attested"] is True
    board = read_json(workspace / "design_board.json")
    assert board["confirmed"] is True
    assert board["confirmed_by"] == "ci"
    assert board["confirm_attested"] is True
    assert board["confirmed_at"]
    assert board["questions"][0]["confirmed"] is True
    # 确认动作本身不改变内容哈希，否则 render→confirm→checker 永远对不上。
    assert board["board_sha256"] == design_board.compute_board_sha256(board)


# ---------------------------------------------------------------------------
# 4. confirm --skip-interactive：留 confirm_not_attested 痕迹
# ---------------------------------------------------------------------------


def test_confirm_skip_interactive_marks_not_attested(tmp_path, monkeypatch):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_ready_workspace(workspace)
    monkeypatch.delenv(design_board.CONFIRM_CODE_ENV, raising=False)

    result = design_board.confirm_board(str(workspace), apply=True, skip_interactive=True, reason="CI 无终端")

    assert result["confirmed"] is True
    assert result["confirmed_by"] == "user_non_interactive"
    assert result["confirm_attested"] is False
    assert result["receipt_origin"] == "skipped_interactive"
    assert result["warning"] == "confirm_not_attested"
    assert result["reason"] == "CI 无终端"

    board = read_json(workspace / "design_board.json")
    assert board["confirmed_by"] == "user_non_interactive"
    assert board["confirm_attested"] is False
    assert board["receipt_origin"] == "skipped_interactive"
    assert board["skip_reason"] == "CI 无终端"

    checked = design_board_checker.check_design_board(str(workspace), "excellence", cross_audit=False)
    assert "confirm_not_attested" in issue_codes(checked)


# ---------------------------------------------------------------------------
# 5. 收敛判据：no_change_rounds=0 时拒绝写入确认字段
# ---------------------------------------------------------------------------


def test_confirm_refuses_when_no_change_rounds_not_converged(tmp_path, monkeypatch):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    board = make_board(no_change_rounds=0)
    board["confirmed"] = False
    write_json(workspace / "design_board.json", board)
    write_rounds(workspace)
    design_board.render_board(str(workspace))
    state = read_json(workspace / "design_board.json")
    monkeypatch.setenv(design_board.CONFIRM_CODE_ENV, state["confirm_code"])

    with pytest.raises(ValueError) as excinfo:
        design_board.confirm_board(str(workspace), apply=True)

    message = str(excinfo.value)
    assert "收敛" in message or "no_change_rounds" in message
    reread = read_json(workspace / "design_board.json")
    assert reread["confirmed"] is False
    assert reread.get("confirm_attested") is not True
    assert reread.get("confirmed_by") is None

    status = design_board.board_status(str(workspace))
    assert status["ready_to_confirm"] is False
    no_change_check = next(check for check in design_board.confirm_board(str(workspace), apply=False)["checks"]
                           if check["name"] == "no_change_rounds")
    assert no_change_check["passed"] is False


# ---------------------------------------------------------------------------
# 6/7. 时间线校验：倒置必须报错，顺序正确不得误报
# ---------------------------------------------------------------------------


def test_checker_flags_code_written_before_confirm(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_confirmed_workspace(workspace, confirmed_at="2099-01-01T00:00:00+00:00")
    touch_at(workspace / "solution.py", "2000-01-01T00:00:00+00:00")

    result = design_board_checker.check_design_board(str(workspace), "excellence", skip_timeline=False)

    assert result["timeline"] == "checked"
    assert result["passed"] is False
    assert "code_written_before_confirm" in issue_codes(result)
    timeline_check = next(check for check in result["checks"] if check["name"] == "timeline")
    assert timeline_check["passed"] is False
    assert [violation["file"] for violation in timeline_check["violations"]] == ["solution.py"]


def test_checker_accepts_code_written_after_confirm(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_confirmed_workspace(workspace, confirmed_at=CONFIRMED_AT)
    touch_at(workspace / "solution.py", "2026-06-28T12:00:10+00:00")

    result = design_board_checker.check_design_board(str(workspace), "excellence", skip_timeline=False)

    assert result["timeline"] == "checked"
    assert "code_written_before_confirm" not in issue_codes(result)
    assert result["passed"] is True


# ---------------------------------------------------------------------------
# 8. 交叉校验：与 model_selection_audit.json 对不上账必须报错
# ---------------------------------------------------------------------------


def test_checker_flags_candidate_mismatch_and_honours_cross_audit_switch(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_confirmed_workspace(workspace, option_names=("A方案", "B方案"))
    write_audit(workspace, ["完全无关的X", "完全无关的Y"], "A方案")

    strict = design_board_checker.check_design_board(str(workspace), "excellence", cross_audit=True)
    assert strict["cross_audit"] == "checked"
    assert "board_audit_candidate_mismatch" in issue_codes(strict)
    assert "board_audit_selection_mismatch" not in issue_codes(strict)
    assert strict["passed"] is False

    rolled_back = design_board_checker.check_design_board(str(workspace), "excellence", cross_audit=False)
    assert rolled_back["cross_audit"] == "skipped"
    assert "board_audit_candidate_mismatch" not in issue_codes(rolled_back)


# ---------------------------------------------------------------------------
# 8'. 完全合法的板必须 passed=True（excellence 模式）
# ---------------------------------------------------------------------------


def test_checker_passes_fully_conformed_board(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    board = seed_confirmed_workspace(workspace)

    result = design_board_checker.check_design_board(str(workspace), "excellence")

    assert result["passed"] is True
    assert result["warning"] is False
    assert result["issues"] == []
    assert result["confirmed"] is True
    assert result["timeline"] == "checked"
    assert result["cross_audit"] == "checked"
    assert len(board["board_sha256"]) == 64


def test_checker_detects_content_tampering_after_confirm(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    seed_confirmed_workspace(workspace)
    board = read_json(workspace / "design_board.json")
    board["questions"][0]["options"][0]["rationale"] = "偷偷改成另一个理由"
    write_json(workspace / "design_board.json", board)

    result = design_board_checker.check_design_board(str(workspace), "excellence")

    assert result["passed"] is False
    assert "board_hash_mismatch" in issue_codes(result)


# ---------------------------------------------------------------------------
# 缺失板的模式语义
# ---------------------------------------------------------------------------


def test_checker_missing_board_warns_in_standard_and_fails_in_excellence(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()

    standard = design_board_checker.check_design_board(str(workspace), "standard")
    assert standard["passed"] is True
    assert standard["warning"] is True
    assert "missing_design_board" in issue_codes(standard)

    excellence = design_board_checker.check_design_board(str(workspace), "excellence")
    assert excellence["passed"] is False
    assert "missing_design_board" in issue_codes(excellence)


# ---------------------------------------------------------------------------
# 哈希口径回归：工具与 checker 必须逐字一致
# ---------------------------------------------------------------------------


def test_agent_and_checker_hash_helpers_agree_and_ignore_confirmation_state(tmp_path):
    board = make_board()
    assert design_board.compute_board_sha256(board) == design_board_checker._content_hash(board)

    mutated = json.loads(json.dumps(board))
    mutated.update(
        {
            "board_sha256": "f" * 64,
            "confirm_code": "AAAA",
            "confirmed": True,
            "confirmed_at": CONFIRMED_AT,
            "confirmed_by": "user",
            "receipt_origin": "browser-submit",
            "receipt_nonce": "nonce-1",
            "confirm_attested": True,
            "no_change_rounds": 2,
            "revision": 9,
            "history": [{"round": 1, "changes": []}],
            "skip_reason": None,
        }
    )
    mutated["questions"][0].update(
        {"user_choice": "O2", "user_comment": "改用备选", "user_locked": True, "confirmed": True}
    )

    assert design_board.compute_board_sha256(mutated) == design_board_checker._content_hash(mutated)
    assert design_board.compute_board_sha256(mutated) == design_board.compute_board_sha256(board)

    # 反过来：真正改动方案内容（候选理由）必须让哈希变化。
    tampered = json.loads(json.dumps(board))
    tampered["questions"][0]["options"][0]["risk"] = "换了一个风险描述"
    assert design_board.compute_board_sha256(tampered) != design_board.compute_board_sha256(board)


# ---------------------------------------------------------------------------
# ingest：回执落盘与收敛轮次
# ---------------------------------------------------------------------------


def test_ingest_submission_writes_round_and_tracks_no_change_rounds(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    board = make_board(user_choice="O1", no_change_rounds=0)
    write_json(workspace / "design_board.json", board)
    submission = workspace / "submission.json"

    write_json(submission, {"decisions": [{"question_id": "Q1", "option_id": "O2"}], "receipt_origin": "browser-submit"})
    first = design_board.ingest_submission(str(workspace), str(submission))
    assert first["round"] == 1
    assert first["user_edited"] is True
    assert first["no_change_rounds"] == 0
    round_one = read_json(workspace / "design_rounds" / "round_001.json")
    assert round_one["user_edited"] is True
    assert round_one["changes"][0]["to"] == "O2"

    write_json(submission, {"decisions": [{"question_id": "Q1", "option_id": "O2"}]})
    second = design_board.ingest_submission(str(workspace), str(submission))
    assert second["round"] == 2
    assert second["user_edited"] is False
    assert second["no_change_rounds"] == 1
    round_two = read_json(workspace / "design_rounds" / "round_002.json")
    assert round_two["user_edited"] is False
    assert round_two["changes"] == []

    after = read_json(workspace / "design_board.json")
    assert after["revision"] == first["revision"] + 1
    assert after["questions"][0]["user_choice"] == "O2"


# ---------------------------------------------------------------------------
# CLI：render / status / confirm --skip-interactive
# ---------------------------------------------------------------------------


def test_cli_render_status_and_skip_interactive_confirm(tmp_path):
    workspace = tmp_path / "ws"
    workspace.mkdir()
    board = make_board()
    board["confirmed"] = False
    write_json(workspace / "design_board.json", board)
    write_rounds(workspace)

    rendered = run_cli("render", "--workspace", str(workspace))
    assert rendered.returncode == 0, rendered.stderr
    render_payload = json.loads(rendered.stdout)
    assert render_payload["confirm_code"] == render_payload["board_sha256"][:4].upper()

    status = run_cli("status", "--workspace", str(workspace))
    assert status.returncode == 0, status.stderr
    status_payload = json.loads(status.stdout)
    assert status_payload["ready_to_confirm"] is True
    assert status_payload["confirm_code"] == render_payload["confirm_code"]
    assert status_payload["confirmed"] is False

    confirmed = run_cli("confirm", "--workspace", str(workspace), "--apply", "--skip-interactive")
    assert confirmed.returncode == 0, confirmed.stderr
    confirm_payload = json.loads(confirmed.stdout)
    assert confirm_payload["warning"] == "confirm_not_attested"
    assert confirm_payload["confirmed_by"] == "user_non_interactive"
    assert read_json(workspace / "design_board.json")["receipt_origin"] == "skipped_interactive"


# ---------------------------------------------------------------------------
# 视图模块：html / xlsx（P1）
# ---------------------------------------------------------------------------


def test_render_all_formats_writes_md_html_xlsx(tmp_path):
    workspace = tmp_path
    seed_board(workspace)

    payload = design_board.render_board(str(workspace), "all")

    assert payload["confirm_code"] == payload["board_sha256"][:4].upper()
    written = [Path(item).name for item in payload["written"]]
    assert "design_board.md" in written
    assert "design_board.html" in written
    assert "design_choices.xlsx" in written
    assert not payload.get("skipped"), payload.get("skipped")

    html = (workspace / "design_board.html").read_text(encoding="utf-8")
    # 离线单文件：不允许任何外链
    for marker in ("http://", "https://", "<link", "<script src", "@import"):
        assert marker not in html, marker
    assert payload["confirm_code"] in html
    assert "<svg" in html  # 内联 SVG 思维导图
    # 每个候选名都要出现在页面上
    board = read_json(workspace / "design_board.json")
    for option in board["questions"][0]["options"]:
        assert option["name"] in html

    from openpyxl import load_workbook

    workbook = load_workbook(workspace / "design_choices.xlsx")
    assert workbook.sheetnames == [
        "00_总览",
        "01_待决策",
        "02_权重",
        "03_打分",
        "04_来源",
        "05_批注",
        "06_说明",
    ]
    decide = workbook["01_待决策"]
    assert decide.protection.sheet is True
    assert decide.data_validations.dataValidation, "01_待决策 必须有下拉数据验证"


def test_render_xlsx_round_trip_through_ingest(tmp_path):
    workspace = tmp_path
    seed_board(workspace)
    design_board.render_board(str(workspace), "xlsx")

    from openpyxl import load_workbook

    xlsx_path = workspace / "design_choices.xlsx"
    workbook = load_workbook(xlsx_path)
    sheet = workbook["01_待决策"]
    headers = [cell.value for cell in sheet[1]]
    choice_col = headers.index("当前选择") + 1
    comment_col = headers.index("批注") + 1
    before = read_json(workspace / "design_board.json")
    current = before["questions"][0]["user_choice"]
    target = next(
        option["id"] for option in before["questions"][0]["options"] if option["id"] != current
    )
    sheet.cell(row=2, column=choice_col).value = target
    sheet.cell(row=2, column=comment_col).value = "Excel 里改的选择"
    workbook.save(xlsx_path)

    result = design_board.ingest_submission(str(workspace), str(xlsx_path))

    assert result["user_edited"] is True
    assert result["changes"][0]["to"] == target
    after = read_json(workspace / "design_board.json")
    assert after["questions"][0]["user_choice"] == target
    # receipt_origin 只落在回执里，板上该字段归 confirm 写
    round_payload = json.loads((workspace / "design_rounds" / Path(result["path"]).name).read_text(encoding="utf-8"))
    assert round_payload["receipt_origin"] == "xlsx-upload"
    assert round_payload["changes"][0]["to"] == target
    # ingest 绝不能写确认字段（板本身不带这些键时也不该被凭空造出来）
    assert after.get("confirmed") is not True
    assert after.get("confirm_attested") is not True


def test_render_xlsx_unchanged_round_trip_is_not_edited(tmp_path):
    workspace = tmp_path
    seed_board(workspace)
    design_board.render_board(str(workspace), "xlsx")

    result = design_board.ingest_submission(str(workspace), str(workspace / "design_choices.xlsx"))

    assert result["user_edited"] is False
    assert result["no_change_rounds"] == 3


def test_html_escapes_injected_markup(tmp_path):
    workspace = tmp_path
    board = seed_board(workspace)
    board["questions"][0]["options"][0]["name"] = "<script>alert(1)</script>"
    write_json(workspace / "design_board.json", board)

    html_module = __import__("tools.design_board_html", fromlist=["render_html"])
    html = html_module.render_html(board)

    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html


def test_html_module_rejects_external_references(tmp_path):
    board = seed_board(tmp_path)
    board["problem_summary"]["key_ambiguities"] = []
    html_module = __import__("tools.design_board_html", fromlist=["render_html"])

    html = html_module.render_html(board)

    for marker in ("http://", "https://", "<link", "<script src"):
        assert marker not in html

# ---------------------------------------------------------------------------
# 反自证：确认码派生、CI/skip 自证、非 TTY 冒充人工
# （这四条对应对抗演练里真实存在的绕过路径，改动这些判据前先看
#   docs/DESIGN_GATE_PLAN.md §20.8 的实测记录）
# ---------------------------------------------------------------------------


def test_checker_flags_hand_forged_confirm_code(tmp_path):
    """手写 confirm_code（不回算 sha 前 4 位）必须被判出来。"""
    board = seed_confirmed_workspace(tmp_path)
    board["confirm_code"] = "ABCD"
    assert board["board_sha256"][:4].upper() != "ABCD"
    write_json(tmp_path / "design_board.json", board)

    strict = design_board_checker.check_design_board(str(tmp_path), "excellence")
    assert "confirm_code_mismatch" in issue_codes(strict)
    assert strict["passed"] is False

    # standard 只降级为 warning，不阻断（兼容旧工作区）
    relaxed = design_board_checker.check_design_board(str(tmp_path), "standard")
    assert "confirm_code_mismatch" in issue_codes(relaxed)


def test_checker_treats_ci_confirmation_as_error_in_excellence(tmp_path):
    """AI 自己用 MM_CONFIRM_CODE 确认会写 confirmed_by=ci，excellence 不能放行。"""
    board = seed_confirmed_workspace(tmp_path)
    board["confirmed_by"] = "ci"
    board["receipt_origin"] = "ci"
    write_json(tmp_path / "design_board.json", board)

    strict = design_board_checker.check_design_board(str(tmp_path), "excellence")
    assert "confirm_ci" in issue_codes(strict)
    assert strict["passed"] is False

    relaxed = design_board_checker.check_design_board(str(tmp_path), "standard")
    assert relaxed["passed"] is True


def test_checker_treats_skip_interactive_as_error_in_excellence(tmp_path):
    """--skip-interactive 写下的 user_non_interactive 在 excellence 下不能当人工确认。"""
    board = seed_confirmed_workspace(tmp_path)
    board["confirmed_by"] = "user_non_interactive"
    board["confirm_attested"] = False
    board["receipt_origin"] = "skipped_interactive"
    write_json(tmp_path / "design_board.json", board)

    strict = design_board_checker.check_design_board(str(tmp_path), "excellence")
    assert "confirm_not_attested" in issue_codes(strict)
    assert strict["passed"] is False

    relaxed = design_board_checker.check_design_board(str(tmp_path), "standard")
    assert relaxed["passed"] is True


def test_confirm_refuses_piped_stdin_without_tty(tmp_path, monkeypatch):
    """`echo <码> | confirm --apply` 不能冒充人工确认（非 TTY 直接拒绝）。"""
    board = seed_ready_workspace(tmp_path)
    monkeypatch.delenv(design_board.CONFIRM_CODE_ENV, raising=False)
    monkeypatch.setattr(design_board, "_is_real_terminal", lambda: False)

    with pytest.raises(ValueError) as excinfo:
        design_board.confirm_board(str(tmp_path), apply=True)

    assert "真实终端" in str(excinfo.value)
    after = read_json(tmp_path / "design_board.json")
    assert after.get("confirmed") is not True
    assert after.get("confirm_attested") is not True


def test_confirm_still_allows_explicit_skip_without_tty(tmp_path, monkeypatch):
    """非 TTY 下 --skip-interactive 仍是可用逃生口（留痕不阻断）。"""
    seed_ready_workspace(tmp_path)
    monkeypatch.delenv(design_board.CONFIRM_CODE_ENV, raising=False)
    monkeypatch.setattr(design_board, "_is_real_terminal", lambda: False)

    result = design_board.confirm_board(str(tmp_path), apply=True, skip_interactive=True, reason="CI")

    assert result["confirmed_by"] == "user_non_interactive"
    assert result["confirm_attested"] is False
